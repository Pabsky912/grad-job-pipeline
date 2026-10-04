#!/usr/bin/env python3
"""grad-job-pipeline helper. Standard library only (Excel export needs openpyxl).

Usage (run from the repository folder):
  python tools/pipeline.py init                      create my-search/ from templates/
  python tools/pipeline.py add --json '{...}'        add or update a tracker row (dedupes)
  python tools/pipeline.py add-contact --json '{...}' add or update a contact row
  python tools/pipeline.py check                     validate tracker.csv and contacts.csv
  python tools/pipeline.py today                     digest: due soon, new finds, follow-ups
  python tools/pipeline.py fetch <ats> <slug> [--q k1,k2] [--json]
                                                     read a public Greenhouse/Lever/Ashby/Workable/Recruitee board
  python tools/pipeline.py dashboard                 build my-search/dashboard.html
  python tools/pipeline.py calendar                  build my-search/deadlines.ics
  python tools/pipeline.py import-changes <file.csv> apply changes exported from the dashboard
  python tools/pipeline.py cv <cv.json> [--out f]    render a one-page printable HTML CV
  python tools/pipeline.py xlsx                      export tracker + contacts to my-search/tracker.xlsx

Options: --workspace <dir> (default: my-search next to this repo, or $GJP_WORKSPACE)
"""
import argparse
import csv
import datetime as dt
import hashlib
import html
import io
import json
import os
import re
import shutil
import sys
import urllib.error
import urllib.request

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..'))
TEMPLATES = os.path.join(ROOT, 'templates')

TRACKER_COLS = ['id', 'track', 'org', 'role', 'location', 'remote', 'link', 'source', 'found_on', 'opens',
                'deadline', 'deadline_note', 'start', 'salary', 'verdict', 'score', 'summary', 'why',
                'watch_out', 'status', 'applied_on', 'next_step', 'next_date', 'last_checked', 'notes']
CONTACT_COLS = ['id', 'name', 'title', 'org', 'channel', 'contact_type', 'why', 'url', 'email',
                'email_verified', 'linked_id', 'template', 'message', 'status', 'sent_on', 'accepted_on',
                'follow_up_on', 'follow_up_stage', 'reply', 'next_step', 'next_date', 'notes']
STATUSES = ['new', 'shortlist', 'this-week', 'later', 'drafting', 'applied', 'interview', 'offer',
            'rejected', 'not-for-me', 'gone']
CLOSED = {'not-for-me', 'gone', 'rejected', 'offer'}
VERDICTS = ['strong', 'worth', 'stretch', 'skip']
CONTACT_STATUSES = ['to-send', 'sent', 'accepted', 'replied', 'meeting', 'no-reply', 'declined', 'closed']
DATE_COLS_T = ['found_on', 'opens', 'deadline', 'applied_on', 'next_date', 'last_checked']
DATE_COLS_C = ['sent_on', 'accepted_on', 'follow_up_on', 'next_date']
UA = 'grad-job-pipeline/1.0 (+https://github.com/Pabsky912/grad-job-pipeline)'


# ----------------------------------------------------------------- helpers
def workspace(args):
    ws = getattr(args, 'workspace', None) or os.environ.get('GJP_WORKSPACE') or os.path.join(ROOT, 'my-search')
    return os.path.abspath(ws)


def today():
    return dt.date.today()


def parse_date(s):
    s = (s or '').strip()
    if not re.fullmatch(r'\d{4}-\d{2}-\d{2}', s):
        return None
    try:
        return dt.date.fromisoformat(s)
    except ValueError:
        return None


def slug(*parts, n=60):
    s = '-'.join(p for p in parts if p)
    s = re.sub(r'[^a-z0-9]+', '-', s.lower()).strip('-')
    return s[:n].strip('-') or 'item'


def read_csv(path, cols):
    if not os.path.exists(path):
        return []
    with open(path, encoding='utf-8-sig', newline='') as f:
        rows = list(csv.DictReader(f))
    return [{c: (r.get(c) or '').strip() for c in cols} | {k: v for k, v in r.items() if k and k not in cols}
            for r in rows]


def write_csv(path, rows, cols):
    extra = []
    for r in rows:
        for k in r:
            if k not in cols and k not in extra:
                extra.append(k)
    tmp = path + '.tmp'
    with open(tmp, 'w', encoding='utf-8', newline='') as f:
        w = csv.DictWriter(f, fieldnames=cols + extra, extrasaction='ignore')
        w.writeheader()
        for r in rows:
            w.writerow({k: r.get(k, '') for k in cols + extra})
    os.replace(tmp, path)


def need_ws(ws):
    if not os.path.isdir(ws):
        sys.exit(f'No workspace at {ws}. Run: python tools/pipeline.py init')


def load_json(path, default):
    try:
        with open(path, encoding='utf-8') as f:
            return json.load(f)
    except (OSError, ValueError):
        return default


def norm(s):
    return re.sub(r'[^a-z0-9]+', ' ', (s or '').lower()).strip()


# ----------------------------------------------------------------- init
def cmd_init(args):
    ws = workspace(args)
    os.makedirs(ws, exist_ok=True)
    for sub in ('cv', 'letters', 'emails'):
        os.makedirs(os.path.join(ws, sub), exist_ok=True)
    made = []
    for name in sorted(os.listdir(TEMPLATES)):
        if '.template' not in name:
            continue
        target = os.path.join(ws, name.replace('.template', ''))
        if os.path.exists(target):
            continue
        shutil.copy2(os.path.join(TEMPLATES, name), target)
        made.append(os.path.basename(target))
    print(f'Workspace: {ws}')
    print('Created: ' + (', '.join(made) if made else 'nothing (all files already exist)'))
    print('Next: the agent runs onboarding (docs/01_ONBOARDING.md).')


# ----------------------------------------------------------------- add
def make_id(row):
    if row.get('id'):
        return slug(row['id'], n=80)
    return slug(row.get('org', ''), row.get('role', ''))


def cmd_add(args):
    ws = workspace(args)
    need_ws(ws)
    path = os.path.join(ws, 'tracker.csv')
    rows = read_csv(path, TRACKER_COLS)
    items = json.loads(args.json)
    items = items if isinstance(items, list) else [items]
    added = updated = 0
    for item in items:
        item = {k: ('' if v is None else ' ; '.join(v) if isinstance(v, list) else str(v)) for k, v in item.items()}
        if not (item.get('org') and item.get('role')):
            print(f'Skipped (needs org and role): {item}')
            continue
        item['id'] = make_id(item)
        item.setdefault('found_on', today().isoformat())
        item.setdefault('status', 'new')
        match = next((r for r in rows if r['id'] == item['id']), None) or next(
            (r for r in rows if norm(r['org']) == norm(item['org']) and norm(r['role']) == norm(item['role'])), None)
        if match:
            for k, v in item.items():
                if k == 'id' or v == '':
                    continue
                if k in ('status', 'found_on') and match.get(k):
                    continue  # never overwrite the person's own status or the first-seen date
                match[k] = v
            updated += 1
        else:
            rows.append({c: item.get(c, '') for c in TRACKER_COLS} | {k: v for k, v in item.items() if k not in TRACKER_COLS})
            added += 1
    write_csv(path, rows, TRACKER_COLS)
    print(f'tracker.csv: {added} added, {updated} updated, {len(rows)} rows in total')


def cmd_add_contact(args):
    ws = workspace(args)
    need_ws(ws)
    path = os.path.join(ws, 'contacts.csv')
    rows = read_csv(path, CONTACT_COLS)
    items = json.loads(args.json)
    items = items if isinstance(items, list) else [items]
    added = updated = 0
    for item in items:
        item = {k: '' if v is None else str(v) for k, v in item.items()}
        if not item.get('name'):
            print(f'Skipped (needs name): {item}')
            continue
        item['id'] = slug(item.get('id') or slug(item['name'], item.get('org', '')))
        item.setdefault('status', 'to-send')
        match = next((r for r in rows if r['id'] == item['id']), None)
        if match:
            match.update({k: v for k, v in item.items() if v != ''})
            updated += 1
        else:
            rows.append({c: item.get(c, '') for c in CONTACT_COLS})
            added += 1
    write_csv(path, rows, CONTACT_COLS)
    print(f'contacts.csv: {added} added, {updated} updated, {len(rows)} rows in total')


# ----------------------------------------------------------------- check
def validate(ws):
    problems = []
    rows = read_csv(os.path.join(ws, 'tracker.csv'), TRACKER_COLS)
    seen, pairs = {}, {}
    for i, r in enumerate(rows, start=2):
        where = f'tracker.csv line {i} ({r.get("org")} / {r.get("role")})'
        if not r['id']:
            problems.append(f'{where}: missing id')
        elif r['id'] in seen:
            problems.append(f'{where}: duplicate id {r["id"]} (also line {seen[r["id"]]})')
        seen.setdefault(r['id'], i)
        key = (norm(r['org']), norm(r['role']))
        if key in pairs and all(key):
            problems.append(f'{where}: same organisation and role as line {pairs[key]}')
        pairs.setdefault(key, i)
        for c in DATE_COLS_T:
            if r[c] and not parse_date(r[c]):
                problems.append(f'{where}: {c} "{r[c]}" is not YYYY-MM-DD')
        if r['status'] and r['status'] not in STATUSES:
            problems.append(f'{where}: unknown status "{r["status"]}" (use one of {", ".join(STATUSES)})')
        if r['verdict'] and r['verdict'] not in VERDICTS:
            problems.append(f'{where}: unknown verdict "{r["verdict"]}"')
        if r['score'] and r['score'] not in list('12345'):
            problems.append(f'{where}: score should be 1-5')
        if r['link'] and not re.match(r'^https?://', r['link']):
            problems.append(f'{where}: link should start with http')
        if not r['link'] and r['status'] not in CLOSED:
            problems.append(f'{where}: no link')
    crow = read_csv(os.path.join(ws, 'contacts.csv'), CONTACT_COLS)
    for i, c in enumerate(crow, start=2):
        where = f'contacts.csv line {i} ({c.get("name")})'
        for col in DATE_COLS_C:
            if c[col] and not parse_date(c[col]):
                problems.append(f'{where}: {col} "{c[col]}" is not YYYY-MM-DD')
        if c['status'] and c['status'] not in CONTACT_STATUSES:
            problems.append(f'{where}: unknown status "{c["status"]}"')
        if c['email'] and c['email_verified'] != 'yes':
            problems.append(f'{where}: email not verified on an official page; check before sending')
        if c['channel'] == 'linkedin' and c['message'] and len(c['message']) > 200 and c['status'] == 'to-send':
            problems.append(f'{where}: connection note is {len(c["message"])} characters (limit 200)')
    return problems, rows, crow


def cmd_check(args):
    ws = workspace(args)
    need_ws(ws)
    problems, rows, crow = validate(ws)
    prefs = load_json(os.path.join(ws, 'preferences.json'), None)
    if prefs is None:
        problems.append('preferences.json is missing or not valid JSON')
    else:
        w = prefs.get('weights') or {}
        if w and abs(sum(float(v) for v in w.values()) - 1) > 0.01:
            problems.append('preferences.json: weights should add up to 1')
        ids = {t.get('id') for t in prefs.get('tracks', [])}
        for r in rows:
            if r['track'] and ids and r['track'] not in ids:
                problems.append(f'tracker.csv: track "{r["track"]}" ({r["org"]}) is not in preferences.json')
                break
    prof = os.path.join(ws, 'profile.md')
    if not os.path.exists(prof) or 'TODO' in open(prof, encoding='utf-8').read():
        problems.append('profile.md still has TODO placeholders: onboarding is not finished')
    print(f'{len(rows)} opportunities, {len(crow)} contacts.')
    print('\n'.join('- ' + p for p in problems) if problems else 'No problems found.')
    return 1 if problems else 0


# ----------------------------------------------------------------- today
def days_left(r):
    d = parse_date(r.get('deadline'))
    return (d - today()).days if d else None


def digest(ws, horizon=14):
    rows = read_csv(os.path.join(ws, 'tracker.csv'), TRACKER_COLS)
    crow = read_csv(os.path.join(ws, 'contacts.csv'), CONTACT_COLS)
    t = today()
    out = []
    due = sorted([r for r in rows if r['status'] not in CLOSED | {'applied', 'interview'}
                  and days_left(r) is not None and 0 <= days_left(r) <= horizon], key=days_left)
    out.append(f'Closing within {horizon} days and not yet applied ({len(due)}):')
    out += [f'  {days_left(r):>3}d  {r["deadline"]}  {r["role"]} · {r["org"]}  [{r["status"] or "new"}]' for r in due] or ['  none']
    opening = [r for r in rows if r['status'] not in CLOSED and parse_date(r['opens'])
               and 0 <= (parse_date(r['opens']) - t).days <= horizon]
    if opening:
        out.append('Opening soon:')
        out += [f'  {r["opens"]}  {r["role"]} · {r["org"]}' for r in opening]
    new = [r for r in rows if r['status'] == 'new' and r['verdict'] in ('strong', 'worth')]
    new.sort(key=lambda r: (VERDICTS.index(r['verdict']), -(int(r['score']) if r['score'].isdigit() else 0)))
    out.append(f'New strong / worth-a-look finds to review ({len(new)}):')
    out += [f'  {r["verdict"]:<6} {r["score"] or "-"}/5  {r["role"]} · {r["org"]}  ({r["deadline"] or r["deadline_note"] or "no deadline"})' for r in new[:10]] or ['  none']
    steps = [x for x in rows + crow if parse_date(x.get('next_date')) and parse_date(x['next_date']) <= t
             and x.get('status') not in CLOSED | {'closed', 'no-reply', 'declined'}]
    out.append(f'Next steps due ({len(steps)}):')
    out += [f'  {x["next_date"]}  {x.get("next_step") or "follow up"} · {x.get("name") or x.get("role")} ({x.get("org")})' for x in steps] or ['  none']
    unverified = [r for r in rows if days_left(r) is not None and 0 <= days_left(r) <= horizon
                  and r['status'] not in CLOSED and (not parse_date(r['last_checked']) or (t - parse_date(r['last_checked'])).days > 7)]
    if unverified:
        out.append(f'Re-check these links (deadline soon, not checked this week): {len(unverified)}')
    return '\n'.join(out)


def cmd_today(args):
    ws = workspace(args)
    need_ws(ws)
    print(f'grad-job-pipeline · {today().strftime("%a %d %b %Y")}')
    print(digest(ws))


# ----------------------------------------------------------------- fetch (public ATS feeds)
def http_json(url, timeout=25):
    req = urllib.request.Request(url, headers={'User-Agent': UA, 'Accept': 'application/json'})
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return json.loads(r.read().decode('utf-8', 'replace'))


def strip_html(s):
    s = re.sub(r'<[^>]+>', ' ', html.unescape(s or ''))
    return re.sub(r'\s+', ' ', html.unescape(s)).strip()


def fetch_board(ats, board):
    ats = ats.lower()
    out = []
    if ats == 'greenhouse':
        data = http_json(f'https://boards-api.greenhouse.io/v1/boards/{board}/jobs?content=true')
        for j in data.get('jobs', []):
            out.append({'id': f'gh-{board}-{j.get("id")}', 'role': j.get('title', ''),
                        'location': (j.get('location') or {}).get('name', ''), 'link': j.get('absolute_url', ''),
                        'posted': (j.get('updated_at') or '')[:10], 'text': strip_html(j.get('content', ''))[:600]})
    elif ats == 'lever':
        data = http_json(f'https://api.lever.co/v0/postings/{board}?mode=json')
        for j in data if isinstance(data, list) else []:
            cat = j.get('categories') or {}
            created = j.get('createdAt')
            out.append({'id': f'lever-{board}-{str(j.get("id", ""))[:8]}', 'role': j.get('text', ''),
                        'location': cat.get('location', ''), 'link': j.get('hostedUrl', ''),
                        'posted': dt.datetime.fromtimestamp(created / 1000, dt.timezone.utc).date().isoformat() if isinstance(created, (int, float)) else '',
                        'text': (j.get('descriptionPlain') or '')[:600], 'team': cat.get('team', ''), 'type': cat.get('commitment', '')})
    elif ats == 'ashby':
        data = http_json(f'https://api.ashbyhq.com/posting-api/job-board/{board}')
        for j in data.get('jobs', []):
            if j.get('isListed') is False:
                continue
            out.append({'id': f'ashby-{board}-{str(j.get("id", ""))[:8]}', 'role': j.get('title', ''),
                        'location': j.get('location', '') + (' (remote)' if j.get('isRemote') else ''),
                        'link': j.get('jobUrl', ''), 'posted': (j.get('publishedAt') or '')[:10],
                        'text': (j.get('descriptionPlain') or '')[:600], 'team': j.get('team', ''), 'type': j.get('employmentType', '')})
    elif ats == 'workable':
        data = http_json(f'https://apply.workable.com/api/v1/widget/accounts/{board}')
        for j in data.get('jobs', []):
            loc = ', '.join(x for x in (j.get('city'), j.get('country')) if x)
            out.append({'id': f'wk-{board}-{j.get("shortcode", "")}', 'role': j.get('title', ''),
                        'location': loc + (' (remote)' if j.get('telecommuting') else ''), 'link': j.get('url', ''),
                        'posted': (j.get('published_on') or '')[:10], 'text': '', 'team': j.get('department', ''),
                        'type': j.get('employment_type', '')})
    elif ats == 'recruitee':
        data = http_json(f'https://{board}.recruitee.com/api/offers/')
        for j in data.get('offers', []):
            out.append({'id': f'rec-{board}-{j.get("id")}', 'role': j.get('title', ''),
                        'location': j.get('location', '') + (' (remote)' if j.get('remote') else ''),
                        'link': j.get('careers_url', ''), 'posted': (j.get('published_at') or '')[:10],
                        'text': strip_html(j.get('description', ''))[:600], 'team': j.get('department', ''),
                        'type': j.get('employment_type_code', '')})
    else:
        raise SystemExit('ats must be one of: greenhouse, lever, ashby, workable, recruitee')
    for o in out:
        o['org_board'] = board
        o['source'] = ats
    return out


def cmd_fetch(args):
    try:
        jobs = fetch_board(args.ats, args.slug)
    except urllib.error.HTTPError as e:
        sys.exit(f'{args.ats}:{args.slug} answered HTTP {e.code}. Check the slug on the company careers page.')
    except (urllib.error.URLError, TimeoutError) as e:
        sys.exit(f'Could not reach {args.ats}:{args.slug} ({e}). No network, or the site is blocked here.')
    total = len(jobs)
    if args.q:
        keys = [k.strip().lower() for k in args.q.split(',') if k.strip()]
        jobs = [j for j in jobs if any(k in (j['role'] + ' ' + j.get('team', '') + ' ' + j['text']).lower() for k in keys)]
    if args.json:
        print(json.dumps(jobs, ensure_ascii=False, indent=1))
        return
    print(f'{args.ats}:{args.slug}: {total} postings, {len(jobs)} match')
    for j in jobs:
        print(f'- {j["role"]} · {j["location"] or "location not stated"} · {j["posted"] or "?"}\n  {j["link"]}  [{j["id"]}]')


# ----------------------------------------------------------------- dashboard
def profile_summary(ws):
    prefs = load_json(os.path.join(ws, 'preferences.json'), {})
    tracks = sorted(prefs.get('tracks', []), key=lambda t: t.get('priority', 99))
    loc = prefs.get('locations', {})
    return {
        'name': prefs.get('name', ''),
        'tracks': [{'id': t.get('id'), 'label': t.get('label') or t.get('id'), 'priority': t.get('priority')} for t in tracks],
        'places': (loc.get('preferred') or []) + (loc.get('acceptable') or []),
        'remote': loc.get('remote_ok'),
        'start': prefs.get('start_from', ''),
        'ruleouts': prefs.get('ruleouts', []),
        'likes': prefs.get('likes', []),
        'dislikes': prefs.get('dislikes', []),
        'targets': prefs.get('weekly_targets', {}),
    }


def cmd_dashboard(args):
    ws = workspace(args)
    need_ws(ws)
    rows = read_csv(os.path.join(ws, 'tracker.csv'), TRACKER_COLS)
    crow = read_csv(os.path.join(ws, 'contacts.csv'), CONTACT_COLS)
    data = {
        'built': dt.datetime.now().strftime('%Y-%m-%d %H:%M'),
        'today': today().isoformat(),
        'profile': profile_summary(ws),
        'items': [{c: r.get(c, '') for c in TRACKER_COLS} for r in rows],
        'contacts': [{c: r.get(c, '') for c in CONTACT_COLS} for r in crow],
        'statuses': STATUSES,
        'contactStatuses': CONTACT_STATUSES,
        'version': hashlib.sha1(json.dumps(rows + crow, sort_keys=True).encode()).hexdigest()[:10],
    }
    tpl = open(os.path.join(HERE, 'dashboard_template.html'), encoding='utf-8').read()
    marker = '/*__DATA__*/null'
    assert marker in tpl, 'dashboard template marker missing'
    payload = json.dumps(data, ensure_ascii=False).replace('</', '<\\/')
    out = os.path.join(ws, 'dashboard.html')
    with open(out, 'w', encoding='utf-8') as f:
        f.write(tpl.replace(marker, payload, 1))
    print(f'Built {out}: {len(rows)} opportunities, {len(crow)} contacts. Open it in any browser.')


# ----------------------------------------------------------------- calendar
def ics_escape(s):
    return (s or '').replace('\\', '\\\\').replace(';', '\\;').replace(',', '\\,').replace('\n', '\\n')


def ics_fold(line):
    b = line.encode('utf-8')
    if len(b) <= 74:
        return line
    parts, cur = [], b''
    for ch in line:
        e = ch.encode('utf-8')
        if len(cur) + len(e) > 73:
            parts.append(cur.decode('utf-8'))
            cur = b''
        cur += e
    parts.append(cur.decode('utf-8'))
    return '\r\n '.join(parts)


def ics_event(uid, day, summary, desc, url='', alarms=(7, 1)):
    stamp = dt.datetime.now(dt.timezone.utc).strftime('%Y%m%dT%H%M%SZ')
    lines = ['BEGIN:VEVENT', f'UID:{uid}@grad-job-pipeline', f'DTSTAMP:{stamp}',
             f'DTSTART;VALUE=DATE:{day.strftime("%Y%m%d")}',
             f'DTEND;VALUE=DATE:{(day + dt.timedelta(days=1)).strftime("%Y%m%d")}',
             f'SUMMARY:{ics_escape(summary)}', f'DESCRIPTION:{ics_escape(desc)}', 'TRANSP:TRANSPARENT']
    if url:
        lines.append(f'URL:{url}')
    for d in alarms:
        lines += ['BEGIN:VALARM', 'ACTION:DISPLAY', f'DESCRIPTION:{ics_escape(summary)}',
                  f'TRIGGER:-P{d}D' if d else 'TRIGGER:PT9H', 'END:VALARM']
    lines.append('END:VEVENT')
    return lines


def cmd_calendar(args):
    ws = workspace(args)
    need_ws(ws)
    rows = read_csv(os.path.join(ws, 'tracker.csv'), TRACKER_COLS)
    crow = read_csv(os.path.join(ws, 'contacts.csv'), CONTACT_COLS)
    ev, n = [], 0
    for r in rows:
        if r['status'] in CLOSED:
            continue
        label = f'{r["role"]} · {r["org"]}'
        desc = '\n'.join(x for x in (r['summary'], f'Verdict: {r["verdict"]} {r["score"]}/5' if r['verdict'] else '',
                                     f'Status: {r["status"]}', r['link']) if x)
        d = parse_date(r['deadline'])
        if d and r['status'] not in ('applied', 'interview'):
            ev += ics_event(f'deadline-{r["id"]}', d, f'Deadline: {label}', desc, r['link']); n += 1
        o = parse_date(r['opens'])
        if o and o >= today():
            ev += ics_event(f'opens-{r["id"]}', o, f'Opens: {label}', desc, r['link'], alarms=(1,)); n += 1
        nd = parse_date(r['next_date'])
        if nd:
            ev += ics_event(f'next-{r["id"]}', nd, f'To do: {r["next_step"] or "next step"} ({r["org"]})', desc, r['link'], alarms=(0,)); n += 1
    for c in crow:
        nd = parse_date(c['next_date'])
        if nd and c['status'] not in ('closed', 'no-reply', 'declined'):
            ev += ics_event(f'contact-{c["id"]}', nd, f'Follow up: {c["name"]} ({c["org"]})',
                            c['next_step'] or c['why'], c['url'], alarms=(0,)); n += 1
    lines = ['BEGIN:VCALENDAR', 'VERSION:2.0', 'PRODID:-//grad-job-pipeline//EN', 'CALSCALE:GREGORIAN',
             'METHOD:PUBLISH', 'X-WR-CALNAME:Job search deadlines'] + ev + ['END:VCALENDAR']
    out = os.path.join(ws, 'deadlines.ics')
    with open(out, 'w', encoding='utf-8', newline='') as f:
        f.write('\r\n'.join(ics_fold(l) for l in lines) + '\r\n')
    print(f'Built {out}: {n} events. Import it into Google Calendar, Outlook or Apple Calendar.')


# ----------------------------------------------------------------- import changes from the dashboard
def cmd_import_changes(args):
    ws = workspace(args)
    need_ws(ws)
    with open(args.file, encoding='utf-8-sig', newline='') as f:
        changes = list(csv.DictReader(f))
    tpath, cpath = os.path.join(ws, 'tracker.csv'), os.path.join(ws, 'contacts.csv')
    rows, crow = read_csv(tpath, TRACKER_COLS), read_csv(cpath, CONTACT_COLS)
    tix, cix = {r['id']: r for r in rows}, {c['id']: c for c in crow}
    done, missing = 0, []
    for ch in changes:
        kind, rid, field, value = ch.get('kind'), ch.get('id'), ch.get('field'), ch.get('value', '')
        if kind == 'item' and field in TRACKER_COLS and field != 'id':
            if rid in tix:
                tix[rid][field] = value; done += 1
            elif ch.get('new_row'):
                row = json.loads(ch['new_row'])
                row['id'] = make_id(row)
                rows.append({c: str(row.get(c, '')) for c in TRACKER_COLS}); tix[row['id']] = rows[-1]; done += 1
            else:
                missing.append(rid)
        elif kind == 'contact' and field in CONTACT_COLS and field != 'id':
            if rid in cix:
                cix[rid][field] = value; done += 1
            else:
                missing.append(rid)
    if not args.dry_run:
        backup = os.path.join(ws, '.backup')
        os.makedirs(backup, exist_ok=True)
        stamp = dt.datetime.now().strftime('%Y%m%d-%H%M%S')
        for p in (tpath, cpath):
            if os.path.exists(p):
                shutil.copy2(p, os.path.join(backup, f'{stamp}-{os.path.basename(p)}'))
        write_csv(tpath, rows, TRACKER_COLS)
        write_csv(cpath, crow, CONTACT_COLS)
    print(f'{done} changes applied' + (' (dry run)' if args.dry_run else '') + (f'; not found: {sorted(set(missing))}' if missing else ''))
    if not args.dry_run:
        cmd_dashboard(args)


# ----------------------------------------------------------------- CV (one-page printable HTML)
CV_CSS = """
@page{size:A4;margin:0}
*{box-sizing:border-box}
body{margin:0;background:#e9e9e6;font-family:Garamond,"EB Garamond","Times New Roman",serif;color:#111}
.page{width:210mm;min-height:297mm;margin:16px auto;background:#fff;padding:9mm 10mm 7mm;box-shadow:0 2px 12px rgba(0,0,0,.15)}
.updated{text-align:right;font-style:italic;color:#666;font-size:8pt}
h1{text-align:center;font-size:18pt;margin:2mm 0 1mm}
.contact{text-align:center;font-size:9pt;margin-bottom:2.5mm}
.contact a{color:#1a4d8f;text-decoration:none}
h2{font-size:12pt;margin:3mm 0 1.2mm;padding-bottom:.8mm;border-bottom:.6pt solid #999}
.row{display:flex;justify-content:space-between;gap:4mm;font-size:9.5pt;line-height:1.25}
.row .r{text-align:right;white-space:nowrap}
.title{font-weight:700}
.sub{font-style:italic}
ul{margin:.6mm 0 2mm;padding-left:4.6mm;list-style:"o  "}
li{font-size:9pt;line-height:1.28;margin:.3mm 0}
.skills p{font-size:9pt;margin:.5mm 0;line-height:1.3}
.sig{text-align:center;font-style:italic;color:#999;font-size:9pt;margin-top:2mm}
@media print{body{background:#fff}.page{margin:0;box-shadow:none}.noprint{display:none}}
.noprint{max-width:210mm;margin:10px auto 0;font:13px system-ui,sans-serif;color:#333}
"""


def rich(parts):
    """A bullet is a string, or a list of [text, bold] pairs."""
    if isinstance(parts, str):
        return html.escape(parts)
    return ''.join(f'<b>{html.escape(t)}</b>' if b else html.escape(t) for t, b in parts)


def cmd_cv(args):
    cv = load_json(args.file, None)
    if not cv:
        sys.exit(f'Could not read {args.file} as JSON (see examples/sample-candidate/cv.json).')
    c = cv.get('contact', {})
    bits = [html.escape(x) for x in (c.get('location'), c.get('email'), c.get('phone')) if x]
    if c.get('link'):
        bits.append(f'<a href="{html.escape(c["link"])}">{html.escape(c.get("link_label") or c["link"])}</a>')
    out = [f'<!doctype html><html lang="en"><head><meta charset="utf-8"><title>CV · {html.escape(cv.get("name", ""))}</title>',
           f'<meta name="viewport" content="width=device-width,initial-scale=1"><style>{CV_CSS}</style></head><body>',
           '<div class="noprint">Print to PDF: A4, margins "None", headers and footers off. Check it fits one page.</div>',
           '<div class="page">']
    if cv.get('last_updated'):
        out.append(f'<div class="updated">Last updated {html.escape(cv["last_updated"])}</div>')
    out.append(f'<h1>{html.escape(cv.get("name", ""))}</h1><div class="contact">{"&nbsp;&nbsp;•&nbsp;&nbsp;".join(bits)}</div>')
    for sec in cv.get('sections', []):
        out.append(f'<h2>{html.escape(sec.get("heading", ""))}</h2>')
        for e in sec.get('entries', []):
            out.append(f'<div class="row"><span class="title">{html.escape(e.get("title", ""))}</span><span class="r title">{html.escape(e.get("where", ""))}</span></div>')
            if e.get('subtitle') or e.get('dates'):
                out.append(f'<div class="row"><span class="sub">{html.escape(e.get("subtitle", ""))}</span><span class="r sub">{html.escape(e.get("dates", ""))}</span></div>')
            if e.get('bullets'):
                out.append('<ul>' + ''.join(f'<li>{rich(b)}</li>' for b in e['bullets']) + '</ul>')
    if cv.get('skills'):
        out.append('<h2>Skills &amp; Languages</h2><div class="skills">' +
                   ''.join(f'<p><b>{html.escape(k)}:</b> {html.escape(v)}</p>' for k, v in cv['skills']) + '</div>')
    out.append(f'<div class="sig">{html.escape(cv.get("name", ""))}</div></div></body></html>')
    target = args.out or os.path.splitext(args.file)[0] + '.html'
    with open(target, 'w', encoding='utf-8') as f:
        f.write('\n'.join(out))
    print(f'Wrote {target}. Open it in a browser and print to PDF; check that it is one page.')


# ----------------------------------------------------------------- xlsx
def cmd_xlsx(args):
    ws = workspace(args)
    need_ws(ws)
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.exit('Excel export needs openpyxl: pip install openpyxl  (tracker.csv already opens in Excel as it is)')
    wb = openpyxl.Workbook()
    head = PatternFill('solid', fgColor='FFDCEAE3')
    for title, fname, cols in (('Opportunities', 'tracker.csv', TRACKER_COLS), ('Contacts', 'contacts.csv', CONTACT_COLS)):
        sh = wb.active if title == 'Opportunities' else wb.create_sheet(title)
        sh.title = title
        rows = read_csv(os.path.join(ws, fname), cols)
        hdr = cols + (['days_left'] if title == 'Opportunities' else [])
        sh.append(hdr)
        for cell in sh[1]:
            cell.font = Font(bold=True); cell.fill = head; cell.alignment = Alignment(wrap_text=True, vertical='top')
        for i, r in enumerate(rows, start=2):
            vals = []
            for c in cols:
                v = r.get(c, '')
                d = parse_date(v) if c in DATE_COLS_T + DATE_COLS_C else None
                vals.append(d if d else (int(v) if c == 'score' and v.isdigit() else v))
            if title == 'Opportunities':
                col = get_column_letter(cols.index('deadline') + 1)
                vals.append(f'=IF({col}{i}="","",{col}{i}-TODAY())')
            sh.append(vals)
            for c in cols:
                cell = sh.cell(i, cols.index(c) + 1)
                if isinstance(cell.value, dt.date):
                    cell.number_format = 'dd mmm yyyy'
                if c in ('link', 'url') and str(cell.value).startswith('http'):
                    cell.hyperlink = cell.value; cell.font = Font(color='FF0563C1', underline='single')
        sh.freeze_panes = 'A2'
        sh.auto_filter.ref = f'A1:{get_column_letter(len(hdr))}{max(2, sh.max_row)}'
        for j, c in enumerate(hdr, start=1):
            sh.column_dimensions[get_column_letter(j)].width = {'summary': 45, 'why': 45, 'watch_out': 40, 'notes': 35,
                                                               'message': 50, 'role': 32, 'org': 24, 'link': 30}.get(c, 14)
    out = os.path.join(ws, 'tracker.xlsx')
    try:
        wb.save(out)
    except PermissionError:
        sys.exit('LOCKED: tracker.xlsx is open in Excel. Close it and run again.')
    print(f'Wrote {out} (tracker.csv stays the master copy).')


# ----------------------------------------------------------------- main
def main(argv=None):
    p = argparse.ArgumentParser(description='grad-job-pipeline helper', formatter_class=argparse.RawDescriptionHelpFormatter,
                                epilog=__doc__.split('\n', 2)[2])
    p.add_argument('--workspace', help='workspace folder (default: my-search/)')
    sub = p.add_subparsers(dest='cmd', required=True)
    sub.add_parser('init').set_defaults(fn=cmd_init)
    a = sub.add_parser('add'); a.add_argument('--json', required=True); a.set_defaults(fn=cmd_add)
    a = sub.add_parser('add-contact'); a.add_argument('--json', required=True); a.set_defaults(fn=cmd_add_contact)
    sub.add_parser('check').set_defaults(fn=cmd_check)
    sub.add_parser('today').set_defaults(fn=cmd_today)
    a = sub.add_parser('fetch'); a.add_argument('ats'); a.add_argument('slug'); a.add_argument('--q', default='')
    a.add_argument('--json', action='store_true'); a.set_defaults(fn=cmd_fetch)
    sub.add_parser('dashboard').set_defaults(fn=cmd_dashboard)
    sub.add_parser('calendar').set_defaults(fn=cmd_calendar)
    a = sub.add_parser('import-changes'); a.add_argument('file'); a.add_argument('--dry-run', action='store_true')
    a.set_defaults(fn=cmd_import_changes)
    a = sub.add_parser('cv'); a.add_argument('file'); a.add_argument('--out'); a.set_defaults(fn=cmd_cv)
    sub.add_parser('xlsx').set_defaults(fn=cmd_xlsx)
    # allow --workspace after the sub-command too
    for sp in sub.choices.values():
        sp.add_argument('--workspace', default=argparse.SUPPRESS, help=argparse.SUPPRESS)
    args = p.parse_args(argv)
    return args.fn(args) or 0


if __name__ == '__main__':
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    sys.exit(main())
